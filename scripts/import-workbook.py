#!/usr/bin/env python3
"""
One-time import of the original "2023 movies.xlsx" workbook.

The workbook is a watchlist, not a pick log: it records which movies exist in
each release year and which ones we saw, but never when we saw them or who
picked them. So this seeds the per-year catalogs and the `seen` flags, and
leaves data/watches.json alone — turn history has to start fresh.

    python3 scripts/import-workbook.py <workbook.xlsx>
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent


def slugify(title, year):
    s = unicodedata.normalize("NFKD", str(title)).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")
    return f"{s}-{year}" if year else s


def norm_title(title):
    """The workbook stores some titles article-last, e.g. 'Brutalist, The'."""
    t = str(title).strip()
    m = re.match(r"^(.*),\s+(The|A|An)$", t)
    return f"{m.group(2)} {m.group(1)}" if m else t


def match_key(title):
    s = unicodedata.normalize("NFKD", norm_title(title)).encode("ascii", "ignore").decode()
    s = re.sub(r"^(the|a|an)\s+", "", s.lower())
    return re.sub(r"[^a-z0-9]+", "", s)


def num(v):
    try:
        f = float(v)
        return int(f) if f == int(f) else f
    except (TypeError, ValueError):
        return None


def read_year_sheet(ws):
    """Columns: Title | Seen | Own | Want to see | Where | IMDb | genre-or-lang."""
    out = []
    for row in ws.iter_rows(min_row=2, max_col=7, values_only=True):
        title = row[0]
        if not title or not str(title).strip():
            continue
        out.append(
            {
                "title": norm_title(title),
                "seen": bool(row[1]),
                "owned": bool(row[2]),
                "wantToSee": bool(row[3]),
                "where": str(row[4]).strip() if row[4] else None,
                "imdb": num(row[5]),
                "genre": str(row[6]).strip() if row[6] and len(str(row[6])) > 3 else None,
            }
        )
    return out


def read_awards_pool(ws):
    """
    The frozen ballot sheets lay movies out in three column-blocks, each a
    title row followed by a row of IMDb / RT critic / RT audience.
    """
    rows = list(ws.iter_rows(max_col=9, values_only=True))
    pool = {}
    for i in range(0, len(rows) - 1, 2):
        titles, scores = rows[i], rows[i + 1]
        for block in (0, 3, 6):
            title = titles[block] if block < len(titles) else None
            if not title or not str(title).strip():
                continue
            vals = [num(scores[block + k]) if block + k < len(scores) else None for k in range(3)]
            pool[match_key(title)] = {
                "title": norm_title(title),
                "imdb": vals[0],
                "rtCritic": vals[1],
                "rtAudience": vals[2],
            }
    return pool


# What each category asks for. The workbook records a category's name and its
# points but never its entry type, so these are set by hand.
#   movie     — pick from the eligible pool
#   person    — an actor, so the movie rows expand to show cast
#   character — a character, so the movie rows expand to show roles
#   movieText — pick a movie, then type the answer; textLabel prompts the field
# `pool: unwatched` flips the list to movies we did NOT see.
#
# Best Animal is movieText rather than character on purpose: cast lists rarely
# name the animals, so free text is the only thing that reliably works.
CATEGORY_TYPES = {
    "best-actor": {"type": "person"},
    "best-actress": {"type": "person"},
    "strangest-role-for-a-familiar-actor": {"type": "person"},
    "favorite-character": {"type": "character"},
    "least-favorite-character": {"type": "character"},
    "best-animal": {"type": "movieText", "textLabel": "The animal"},
    "best-scene": {"type": "movieText", "textLabel": "The scene"},
    "biggest-twist": {"type": "movieText", "textLabel": "The twist"},
    "best-original-song": {"type": "movieText", "textLabel": "The song"},
    "most-thought-we-d-seebut-didn-t": {"type": "movie", "pool": "unwatched"},
}


def read_categories(ws):
    """Column A is the category, B the winner's points, C-F honorable mentions."""
    cats = []
    for row in ws.iter_rows(min_row=2, max_col=6, values_only=True):
        name = row[0]
        if not name or not str(name).strip():
            continue
        winner = num(row[1])
        mentions = [num(v) for v in row[2:6] if num(v) is not None]
        cid = slugify(name, None)
        spec = CATEGORY_TYPES.get(cid, {})
        entry = {
            "id": cid,
            "name": str(name).strip(),
            "type": spec.get("type", "movie"),
            "pool": spec.get("pool", "watched"),
            "winnerPoints": winner,
            "honorableMentionPoints": mentions,
            # Unscored categories are still voted on and still appear in the
            # reveal; they just contribute nothing to the totals.
            "scored": winner is not None,
            "allowWriteIn": True,
        }
        if "textLabel" in spec:
            entry["textLabel"] = spec["textLabel"]
        cats.append(entry)
    return cats


def main():
    src = Path(sys.argv[1])
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)

    year_sheets = sorted(
        (n for n in wb.sheetnames if re.fullmatch(r"20\d\d", n)), reverse=True
    )
    frozen_pools = {
        int(n.split()[0]): read_awards_pool(wb[n])
        for n in wb.sheetnames
        if re.fullmatch(r"20\d\d Awards", n)
    }

    (ROOT / "data" / "catalog").mkdir(parents=True, exist_ok=True)
    summary = []

    for name in year_sheets:
        year = int(name)
        entries = read_year_sheet(wb[name])
        if not entries:
            continue
        pool = frozen_pools.get(year, {})

        catalog = []
        for e in entries:
            key = match_key(e["title"])
            extra = pool.get(key, {})
            catalog.append(
                {
                    "id": slugify(e["title"], year),
                    "title": e["title"],
                    "kind": "streaming" if e["where"] else "theatrical",
                    "festivalDate": None,
                    "usLimitedDate": None,
                    "usTheatricalDate": None,
                    "homeDate": None,
                    "isForeignLanguage": False,
                    "isDocumentary": bool(e["genre"] and "document" in e["genre"].lower()),
                    "hadUSTheatricalRelease": not e["where"],
                    "services": [e["where"]] if e["where"] else [],
                    "oscarNominated": False,
                    "ratings": {
                        "imdb": extra.get("imdb", e["imdb"]),
                        "metacritic": None,
                        "rtCritic": extra.get("rtCritic"),
                        "rtAudience": extra.get("rtAudience"),
                    },
                    "cast": [],
                    "imdbId": None,
                    "tmdbId": None,
                    "sources": ["workbook"],
                    # Imported rows have no release dates, so the rules engine
                    # cannot re-derive the year. The sheet it came from is the
                    # authority until a catalog refresh fills the dates in.
                    "computedYear": year,
                    "confidence": "medium",
                    "evidence": [f"Imported from the {year} sheet of the original workbook."],
                    "manual": True,
                    "seen": e["seen"],
                    "owned": e["owned"],
                    "wantToSee": e["wantToSee"],
                    "onFrozenBallot": key in pool,
                }
            )

        path = ROOT / "data" / "catalog" / f"{year}.json"
        path.write_text(json.dumps(catalog, indent=2) + "\n")
        summary.append(
            {
                "year": year,
                "titles": len(catalog),
                "seen": sum(1 for c in catalog if c["seen"]),
                "frozen": len(pool),
            }
        )

    # A lightweight index for title autocomplete. The full catalogs are ~5 MB
    # in total, far too much to pull into the browser just to search titles.
    index = []
    for s in summary:
        year = s["year"]
        for m in json.loads((ROOT / "data" / "catalog" / f"{year}.json").read_text()):
            index.append([m["id"], m["title"], year, 1 if m["seen"] else 0])
    index.sort(key=lambda r: (-r[2], r[1]))
    (ROOT / "data" / "index.json").write_text(json.dumps(index, separators=(",", ":")) + "\n")
    print(f"  index: {len(index)} titles")

    cats = read_categories(wb["Awards"])
    (ROOT / "data" / "categories.json").write_text(json.dumps(cats, indent=2) + "\n")

    for s in summary:
        note = f"  ({s['frozen']} on the frozen ballot)" if s["frozen"] else ""
        print(f"  {s['year']}: {s['titles']:4} titles, {s['seen']:3} seen{note}")
    print(f"\n  {len(cats)} award categories")
    print(f"  {sum(s['seen'] for s in summary)} movies seen in total")


if __name__ == "__main__":
    main()
